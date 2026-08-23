"""UI 元素统一查询入口，只读取项目内本地 Excel。"""

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from playwright.sync_api import Locator, Page

REGISTRY_FILE = Path(__file__).with_name("ui_elements.xlsx")


@dataclass(frozen=True)
class ElementDefinition:
    """本地 Excel 中的一条元素定义。"""

    element_id: int
    project_name: str
    module_name: str
    page_name: str
    name: str
    locating_method: str | int
    expression: str
    index: int | None = None
    tag: str | None = None
    input_type: str | None = None
    role: str | None = None
    description: str | None = None
    interactive: bool = False
    disabled: bool = False

    @classmethod
    def from_record(cls, record: dict[str, Any]) -> "ElementDefinition":
        return cls(
            element_id=int(record["ID"]),
            project_name=str(record["项目名称"]),
            module_name=str(record["模块名称"]),
            page_name=str(record["页面名称"]),
            name=str(record["元素名称"]),
            locating_method=record["定位方式1"],
            expression=str(record["表达式1"]),
            index=_optional_int(record.get("下标1")),
            tag=_optional_text(record.get("标签")),
            input_type=_optional_text(record.get("input类型")),
            role=_optional_text(record.get("role")),
            description=_optional_text(record.get("说明")),
            interactive=_bool(record.get("可交互")),
            disabled=_bool(record.get("禁用")),
        )

    def locate(self, page: Page) -> Locator:
        """按照统一定义生成 Playwright Locator。"""
        method = str(self.locating_method).strip().upper()
        if method in {"TEST_ID", "1"}:
            locator = page.get_by_test_id(self.expression)
        elif method in {"CSS", "9"}:
            locator = page.locator(self.expression)
        elif method in {"TEXT", "3"}:
            locator = page.get_by_text(self.expression, exact=True)
        elif method in {"PLACEHOLDER", "4"}:
            locator = page.get_by_placeholder(self.expression)
        elif method in {"XPATH", "0"}:
            expression = self.expression
            locator = page.locator(expression if expression.startswith("xpath=") else f"xpath={expression}")
        else:
            raise ValueError(f"元素 {self.name} 使用了不支持的定位方式：{self.locating_method}")
        return locator.nth(self.index) if self.index is not None else locator


class ElementRepository:
    """从本地 Excel 按元素引用查询唯一元素定义。"""

    source = "excel"

    @lru_cache(maxsize=512)
    def get(self, name: str, module_name: str | None = None, page_name: str | None = None) -> ElementDefinition:
        """查询唯一元素；同名时必须补充模块或页面。"""
        return self._get_from_local(name, module_name, page_name)

    def locate(
        self,
        page: Page,
        name: str,
        module_name: str | None = None,
        page_name: str | None = None,
    ) -> Locator:
        return self.get(name, module_name, page_name).locate(page)

    @staticmethod
    @lru_cache(maxsize=1)
    def all_local() -> tuple[ElementDefinition, ...]:
        if not REGISTRY_FILE.is_file():
            raise FileNotFoundError(f"UI 元素 Excel 不存在：{REGISTRY_FILE}")
        workbook = load_workbook(REGISTRY_FILE, read_only=True, data_only=True)
        try:
            sheet = workbook["UI元素"]
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() for value in next(rows)]
            definitions = tuple(
                ElementDefinition.from_record(dict(zip(headers, values)))
                for values in rows
                if any(value not in (None, "") for value in values)
            )
        finally:
            workbook.close()
        names = [definition.name for definition in definitions]
        if len(definitions) != 260 or len(names) != len(set(names)):
            raise ValueError("UI Mock 本地元素 Excel 必须包含 260 个唯一元素")
        return definitions

    def _get_from_local(
        self,
        name: str,
        module_name: str | None,
        page_name: str | None,
    ) -> ElementDefinition:
        matched = [
            definition
            for definition in self.all_local()
            if definition.name == name
            and (module_name is None or definition.module_name == module_name)
            and (page_name is None or definition.page_name == page_name)
        ]
        return _unique(name, matched)

def _unique(name: str, matched: list[ElementDefinition]) -> ElementDefinition:
    if not matched:
        raise KeyError(f"UI 元素仓库中不存在：{name}")
    if len(matched) > 1:
        raise ValueError(f"UI 元素仓库中存在 {len(matched)} 个同名元素，请指定模块或页面：{name}")
    return matched[0]


def _optional_int(value: Any) -> int | None:
    return None if value in (None, "") else int(value)


def _optional_text(value: Any) -> str | None:
    return None if value in (None, "") else str(value)


def _bool(value: Any) -> bool:
    return str(value).strip().lower() in {"是", "true", "1", "yes"}


elements = ElementRepository()
