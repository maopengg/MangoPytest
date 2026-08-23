"""从项目内 Excel 加载 UI 自动化用例。"""

import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .models import ElementCase, InventoryCase, OperationCase


CASE_FILE = Path(__file__).resolve().parents[1] / "data" / "MangoMock_UI全元素操作测试用例.xlsx"


def _rows(sheet_name: str) -> Iterable[dict[str, Any]]:
    if not CASE_FILE.is_file():
        raise FileNotFoundError(f"UI Excel 用例文件不存在：{CASE_FILE}")
    workbook = load_workbook(CASE_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() for value in next(iterator)]
        for values in iterator:
            row = dict(zip(headers, values))
            if any(value not in (None, "") for value in values):
                yield row
    finally:
        workbook.close()


def _json(value: Any) -> dict[str, Any]:
    if value in (None, ""):
        return {}
    if isinstance(value, dict):
        return value
    parsed = json.loads(str(value))
    if not isinstance(parsed, dict):
        raise ValueError(f"Excel 操作参数必须是 JSON Object，实际为：{value}")
    return parsed


def _yes(value: Any) -> bool:
    return str(value).strip().lower() in {"是", "true", "1", "yes"}


def _load_operation_cases() -> tuple[OperationCase, ...]:
    cases = tuple(
        OperationCase(
            case_id=str(row["用例ID"]),
            group=str(row["分组"]),
            method=str(row["方法"]),
            description=str(row["方法说明"]),
            target_id=str(row["目标元素ID"]),
            params=_json(row["操作参数(JSON)"]),
            expected=str(row["预期结果"]),
            verify_id=str(row["验证元素ID"]),
            priority=str(row["优先级"]),
        )
        for row in _rows("操作方法覆盖")
    )
    _validate(cases, expected_count=103, label="操作方法")
    return cases


def _load_element_cases() -> tuple[ElementCase, ...]:
    cases = tuple(
        ElementCase(
            case_id=str(row["用例ID"]),
            page_key=str(row["所属页面"]),
            element_id=str(row["元素ID"]),
            control_type=str(row["控件类型"]),
            method=str(row["执行方法"]),
            params=_json(row["参数(JSON)"]),
            expected=str(row["预期结果"]),
            risk=str(row["风险级别"]),
        )
        for row in _rows("交互元素逐一操作")
    )
    _validate(cases, expected_count=130, label="交互元素")
    return cases


def _load_inventory_cases() -> tuple[InventoryCase, ...]:
    cases = tuple(
        InventoryCase(
            sequence=int(row["序号"]),
            page_key=str(row["所属页面"]),
            element_id=str(row["元素ID"]),
            tag=str(row["标签"]),
            interactive=_yes(row["可交互"]),
            disabled=_yes(row["禁用"]),
            strategy=str(row["测试策略"]),
        )
        for row in _rows("全元素清单")
    )
    _validate(cases, expected_count=260, label="全元素")
    return cases


def _validate(cases: tuple[Any, ...], expected_count: int, label: str) -> None:
    case_ids = [case.case_id for case in cases]
    if len(cases) != expected_count:
        raise ValueError(f"{label} Excel 用例数量错误：期望 {expected_count}，实际 {len(cases)}")
    if len(case_ids) != len(set(case_ids)):
        raise ValueError(f"{label} Excel 用例 ID 存在重复")


OPERATION_CASES = _load_operation_cases()
ELEMENT_CASES = _load_element_cases()
INVENTORY_CASES = _load_inventory_cases()

