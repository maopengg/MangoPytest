"""读取并校验 pytest_ui 自己维护的本地 Excel 元素工作簿。"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

WORKBOOK_DIR = Path(__file__).with_name("workbooks")
SHEET_NAME = "UI元素"
REQUIRED_HEADERS = {
    "ID",
    "项目名称",
    "模块名称",
    "页面名称",
    "元素名称",
    "定位方式1",
    "表达式1",
}


def load_element_records() -> tuple[dict[str, Any], ...]:
    files = tuple(sorted(WORKBOOK_DIR.glob("*.xlsx")))
    if not files:
        raise FileNotFoundError(f"未找到本地 UI 元素 Excel：{WORKBOOK_DIR}")
    records = tuple(record for file in files for record in _read_workbook(file))
    _validate_records(records)
    return records


def _read_workbook(file: Path) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(file, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"元素工作簿缺少工作表 {SHEET_NAME}：{file}")
        rows = workbook[SHEET_NAME].iter_rows(values_only=True)
        headers = tuple(str(value).strip() for value in next(rows))
        missing = REQUIRED_HEADERS.difference(headers)
        if missing:
            raise ValueError(f"元素工作簿缺少字段 {sorted(missing)}：{file}")
        for values in rows:
            if any(value not in (None, "") for value in values):
                yield dict(zip(headers, values))
    finally:
        workbook.close()


def _validate_records(records: tuple[dict[str, Any], ...]) -> None:
    if not records:
        raise ValueError("pytest_ui 本地元素 Excel 不能为空")
    names = [str(record["元素名称"]).strip() for record in records]
    ids = [int(record["ID"]) for record in records]
    if len(names) != len(set(names)):
        raise ValueError("pytest_ui 本地元素 Excel 存在重复元素名称")
    if len(ids) != len(set(ids)):
        raise ValueError("pytest_ui 本地元素 Excel 存在重复元素 ID")

