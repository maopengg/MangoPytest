"""通用 Excel 只读数据源。

业务项目只负责把记录转换为自己的模型，不再直接依赖 openpyxl。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


class ExcelWorkbookSource:
    """按工作表读取表头化记录。"""

    def __init__(self, file: str | Path):
        self.file = Path(file)

    def records(
        self,
        sheet_name: str,
        *,
        required_headers: Iterable[str] = (),
    ) -> tuple[dict[str, Any], ...]:
        if not self.file.is_file():
            raise FileNotFoundError(f"Excel 文件不存在：{self.file}")
        workbook = load_workbook(self.file, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Excel 缺少工作表 {sheet_name}：{self.file}")
            rows = workbook[sheet_name].iter_rows(values_only=True)
            try:
                headers = tuple(str(value).strip() for value in next(rows))
            except StopIteration as exc:
                raise ValueError(f"Excel 工作表为空：{self.file} / {sheet_name}") from exc
            missing = set(required_headers).difference(headers)
            if missing:
                raise ValueError(f"Excel 缺少字段 {sorted(missing)}：{self.file}")
            return tuple(
                dict(zip(headers, values))
                for values in rows
                if any(value not in (None, "") for value in values)
            )
        finally:
            workbook.close()

